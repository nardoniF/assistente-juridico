from __future__ import annotations

import html as html_lib
import re
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


_HEADING_RE = re.compile(
    r"^[A-ZÁÉÍÓÚÃÕÇ0-9][A-ZÁÉÍÓÚÃÕÇ0-9 \-–—\.,:;()]{10,}$"
)
_INLINE_FMT = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|_[^_]+_)")


def _strip_md_noise(line: str) -> str:
    s = line.strip()
    s = re.sub(r"^#{1,6}\s+", "", s)
    # Não remover * do meio/fim: **negrito** precisa chegar intacto.
    return s.strip()


def _is_heading(line: str) -> bool:
    raw = line.strip()
    if raw.startswith(("# ", "## ", "### ")):
        return True
    body = _strip_md_noise(raw)
    if len(body) < 8:
        return False
    if body.endswith(":") and len(body) < 80:
        return True
    return bool(_HEADING_RE.match(body))


def _iter_blocks(text: str, title: str | None = None):
    if title:
        yield True, title.strip()
    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        heading = _is_heading(line)
        body = _strip_md_noise(line)
        if not body:
            continue
        yield heading, body


def _add_runs_with_inline(paragraph, text: str, *, bold_all: bool = False, size: Pt = Pt(12)):
    """Aplica **negrito** e *itálico* simples no Word."""
    parts = _INLINE_FMT.split(text)
    if len(parts) == 1 and not bold_all:
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = size
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        return
    for part in parts:
        if not part:
            continue
        bold = bold_all
        italic = False
        chunk = part
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            chunk = part[2:-2]
            bold = True
        elif (part.startswith("*") and part.endswith("*") and len(part) > 2) or (
            part.startswith("_") and part.endswith("_") and len(part) > 2
        ):
            chunk = part[1:-1]
            italic = True
        run = paragraph.add_run(chunk)
        run.bold = bold
        run.italic = italic
        run.font.name = "Times New Roman"
        run.font.size = size
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def markdown_to_docx(text: str, dest: Path, title: str | None = None) -> Path:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(3)
        s.right_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(8)
    pf.space_before = Pt(0)

    title_done = False
    for heading, body in _iter_blocks(text, title=title):
        if title and heading and body == title.strip() and not title_done:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(18)
            p.paragraph_format.first_line_indent = Cm(0)
            _add_runs_with_inline(p, body, bold_all=True, size=Pt(14))
            title_done = True
            continue
        p = doc.add_paragraph()
        if heading:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(10)
            _add_runs_with_inline(p, body, bold_all=True, size=Pt(12))
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.25)
            _add_runs_with_inline(p, body, size=Pt(12))

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return dest


def _inline_html(text: str) -> str:
    parts = _INLINE_FMT.split(text)
    out: list[str] = []
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append(f"<b>{html_lib.escape(part[2:-2])}</b>")
        elif (part.startswith("*") and part.endswith("*") and len(part) > 2) or (
            part.startswith("_") and part.endswith("_") and len(part) > 2
        ):
            out.append(f"<i>{html_lib.escape(part[1:-1])}</i>")
        else:
            out.append(html_lib.escape(part))
    return "".join(out)


def _text_to_story_html(text: str, title: str | None = None) -> str:
    chunks: list[str] = ['<div id="peca">']
    title_done = False
    for heading, body in _iter_blocks(text, title=title):
        if title and heading and body == title.strip() and not title_done:
            chunks.append(f'<p class="titulo">{_inline_html(body)}</p>')
            title_done = True
            continue
        if heading:
            chunks.append(f'<p class="secao">{_inline_html(body)}</p>')
        elif re.match(r"^(\d+[\.\)]\s+|[-•]\s+)", body):
            chunks.append(f'<p class="item">{_inline_html(body)}</p>')
        else:
            chunks.append(f"<p>{_inline_html(body)}</p>")
    chunks.append("</div>")
    return "\n".join(chunks)


_PDF_CSS = """
@page {
  size: A4;
  margin: 2.5cm 2cm 2.2cm 3cm;
}
body {
  font-family: times;
  font-size: 12pt;
  color: #111;
}
#peca p {
  margin: 0 0 10pt 0;
  line-height: 1.5;
  text-align: justify;
  text-indent: 1.25cm;
}
#peca p.titulo {
  text-align: center;
  text-indent: 0;
  font-size: 14pt;
  font-weight: bold;
  margin: 0 0 18pt 0;
  line-height: 1.35;
}
#peca p.secao {
  text-align: center;
  text-indent: 0;
  font-weight: bold;
  font-size: 12pt;
  margin: 14pt 0 10pt 0;
  line-height: 1.35;
}
#peca p.item {
  text-indent: 0;
  padding-left: 0.6cm;
  margin: 0 0 8pt 0;
}
"""


def markdown_to_pdf(text: str, dest: Path, title: str | None = None) -> Path:
    """PDF forense A4: justificado, margens padrão, seções e numeração."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    story_html = _text_to_story_html(text, title=title)
    story = fitz.Story(html=story_html, user_css=_PDF_CSS)

    mediabox = fitz.paper_rect("a4")
    # Margens alinhadas ao CSS (Story place usa o retângulo útil)
    where = mediabox + (85, 70, -56, -62)  # L T R B em pontos (aprox. 3/2.5/2/2.2 cm)

    writer = fitz.DocumentWriter(str(dest))
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()

    # Numeração de páginas (rodapé central)
    doc = fitz.open(dest)
    total = doc.page_count
    for i, page in enumerate(doc, start=1):
        label = f"{i}"
        if total > 1:
            label = f"{i} / {total}"
        page.insert_text(
            (mediabox.width / 2 - 12, mediabox.height - 36),
            label,
            fontname="times-roman",
            fontsize=9,
            color=(0.25, 0.25, 0.25),
        )
    doc.saveIncr()
    doc.close()
    return dest


def save_peca(text: str, case_dir: Path, base_name: str, title: str | None = None) -> dict:
    """Grava Word (editar) + PDF (enviar) na pasta do processo."""
    base = base_name
    if base.lower().endswith(".docx"):
        base = base[:-5]
    elif base.lower().endswith(".pdf"):
        base = base[:-4]
    docx_path = case_dir / f"{base}.docx"
    pdf_path = case_dir / f"{base}.pdf"
    markdown_to_docx(text, docx_path, title=title)
    markdown_to_pdf(text, pdf_path, title=title)
    return {"docx": docx_path.name, "pdf": pdf_path.name, "base": base}
